"""The quotation written into the customer's own spreadsheet layout.

Some customers want the offer back in the file shape they sent the request in,
not on our letterhead. This is that shape: the template made by
`src.tools.make_customer_template`, one row per line, in the customer's order.

Row 2 of the template is the model row. Each item row takes its look from it -
fill, borders, font, the text format on the code column - so the look is the
template's to decide and this module only knows which value goes in which
column.

The numbers are the approved ones and are written as numbers, so the customer
can sum or sort the column without retyping it. The arithmetic is the same
`QuotedLine.total` the PDF uses: one quotation, one rule for its totals.
"""

from collections.abc import Sequence
from copy import copy
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.services.quotation import QuotedLine, number_of

FIRST_ROW = 2
# Category, SN, Code, Description, Unit, Qty, Unit Prc, Item Total.
COLUMNS = 8


def fill(template: bytes, lines: Sequence[QuotedLine]) -> bytes:
    """A copy of the template with one row per line."""
    book = load_workbook(BytesIO(template))
    sheet = book.active
    model = [
        copy(sheet.cell(row=FIRST_ROW, column=column)._style) for column in range(1, COLUMNS + 1)
    ]

    for offset, line in enumerate(lines):
        _write(sheet, FIRST_ROW + offset, line, model)
    if not lines:
        # An empty styled row reads as a line somebody forgot to fill in.
        sheet.delete_rows(FIRST_ROW)

    out = BytesIO()
    book.save(out)
    return out.getvalue()


def _write(sheet: Worksheet, row: int, line: QuotedLine, model: list) -> None:
    quantity = number_of(line.quantity)
    total = line.total
    values = [
        # We have no category for a line, and the column stays blank rather
        # than guessed: the customer's own grouping is theirs to fill in.
        None,
        line.number,
        line.code or None,
        line.description or None,
        line.uom or None,
        # The customer's own words when they are not a number ("2 coil"),
        # rather than a blank that hides what was asked.
        float(quantity) if quantity is not None else (line.quantity or None),
        float(line.unit_price) if line.unit_price is not None else None,
        float(total) if total is not None else None,
    ]
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell._style = copy(model[column - 1])
