"""Build the blank customer-file template the quotation is written into.

    uv run python -m src.tools.make_customer_template

The customer's own layout, as `docs/samples/SSG POC - Supported Customer RFQ
Template Example.xlsx` has it: one sheet, a header row, one row per item. The
blank is drawn afresh rather than cleared out of that example: the example is a
Google Sheets export and carries its authors along (`xl/persons`, a drawing),
and a template shipped in the image must not put their names on every file the
desk sends.

Row 2 is left in, empty but styled: it is the model every item row is copied
from, so the look of the rows lives in the template and not in the code that
fills it. Change the template, and the next file follows.
"""

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side

SHEET = "Customer RFQ Template"
HEADINGS = ("Category", "SN", "Code", "Description", "Unit", "Qty", "Unit Prc", "Item Total")
# Measured off the example, in Excel's own character units.
WIDTHS = {"A": 11.38, "B": 14.5, "C": 14.5, "D": 70.13, "E": 14.5, "F": 14.5, "G": 14.5, "H": 14.5}
ROW_HEIGHT = 15.75

HEADING_FILL = "FFCCF2F4"
ROW_FILL = "FFE0F9F9"
# The column the customer's code goes in. Text, not a number: `04361753` is a
# code, and Excel would drop its leading zero the moment it read it as one.
CODE_COLUMN = "C"


def build(out: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.sheet_format.defaultRowHeight = ROW_HEIGHT
    sheet.sheet_format.customHeight = True

    thin = Side(style="thin", color="FF000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ink = Color(theme=1)

    for column, heading in enumerate(HEADINGS, start=1):
        cell = sheet.cell(row=1, column=column, value=heading)
        cell.font = Font(name="Arial", size=10, bold=True, color=ink)
        cell.fill = PatternFill("solid", fgColor=HEADING_FILL, bgColor=HEADING_FILL)
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

        model = sheet.cell(row=2, column=column)
        model.font = Font(name="Arial", size=10, color=ink)
        model.fill = PatternFill("solid", fgColor=ROW_FILL, bgColor=ROW_FILL)
        model.border = border
        if model.column_letter == CODE_COLUMN:
            model.quotePrefix = True
            model.number_format = "@"

    for letter, width in WIDTHS.items():
        sheet.column_dimensions[letter].width = width

    # Nobody's name on a file that goes to customers.
    book.properties.creator = "Seven Seas Group"
    book.properties.lastModifiedBy = None
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--out",
        type=Path,
        default=Path("config/quotation_customer_file_template.xlsx"),
        help="where to write it",
    )
    args = parser.parse_args()
    build(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
