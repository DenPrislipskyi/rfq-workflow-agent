"""The endpoint the Quote Overview page reads.

The contract under test is not "does it return rows". It is that a record holds
the customer's own email, the model's reasoning and every file they attached,
and that **none of that** is on the wire: the page gets the label, the status
and who wrote in, and a field nobody put there on purpose is a bug.
"""

import asyncio
from contextlib import suppress
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from src import register_exceptions, register_routers
from src.api import quotes
from src.api.dependencies import (
    get_catalog,
    get_changes,
    get_customer_file_template,
    get_quotation_logo,
    get_quote_template,
    get_records,
)
from src.domain.enums import (
    DecisionPath,
    DeliveryOutcome,
    Direction,
    EmailCategory,
    Priority,
    RecommendedAction,
)
from src.domain.models import (
    Attachment,
    ClassificationOutcome,
    ClassificationResult,
    EmailAddress,
    Hints,
    NormalizedEmail,
    Signals,
    SplitThread,
)
from src.domain.rules.catalog import Catalog
from src.infrastructure.documents.loader import SourceFile
from src.infrastructure.storage.changes import Changes
from src.infrastructure.storage.records import (
    EmailRecords,
    RecordedCandidate,
    RecordedDelivery,
    RecordedMatch,
)

URL = "/api/v1/quotes"
CUSTOMER_TEMPLATE = Path(__file__).parents[1] / "config" / "quotation_customer_file_template.xlsx"
QUOTE_TEMPLATE = Path(__file__).parents[1] / "config" / "quotation_sg_uae_workbook_template.xlsm"
BODY = "Dear Sir/Madam, you may find attached our RFQ for Engine Materials."


# The three products the confirmation tests are allowed to settle on. A code
# outside this sheet is a code nobody sells, and the endpoint has to refuse it.
#
# `Supplier` and `Price` are here because they are what the sheet is edited
# for: they are the columns that move between the day a line was matched and
# the day somebody settles it.
SHEET = [
    {
        "Item Code": "T69128400",
        "Item Description": "HEX HEAD BOLT/NUT, M16 X 65MM",
        "Price": "0.51",
        "Supplier": "Northgate Marine Fasteners Ltd.",
    },
    {
        "Item Code": "T69133100",
        "Item Description": "HEX HEAD BOLT/NUT, M20 X 80MM",
        "Price": "0.88",
        "Supplier": "Northgate Marine Fasteners Ltd.",
    },
    {
        "Item Code": "T85116300",
        "Item Description": "WELDER GLOVES FIVE FINGERS",
        "Price": "4.10",
        "Supplier": "Harbour Safety Equipment Co.",
    },
]


def catalog() -> SimpleNamespace:
    """Enough of `CatalogService` for the endpoint: the sheet it reads."""
    return SimpleNamespace(
        current=Catalog.from_rows(
            SHEET, code_column="Item Code", description_column="Item Description"
        )
    )


def client(records: EmailRecords, changes: Changes | None = None) -> TestClient:
    """The app without its lifespan: no Graph subscription, no `.env`."""
    signal = changes or Changes()
    app = FastAPI()
    register_exceptions(app)
    register_routers(app)
    app.dependency_overrides[get_records] = lambda: records
    app.dependency_overrides[get_changes] = lambda: signal
    app.dependency_overrides[get_catalog] = catalog
    app.dependency_overrides[get_quotation_logo] = lambda: None
    app.dependency_overrides[get_customer_file_template] = lambda: CUSTOMER_TEMPLATE
    app.dependency_overrides[get_quote_template] = lambda: QUOTE_TEMPLATE
    return TestClient(app)


def email(sender: str = "purchasing@newcompany.example.com", **overrides) -> NormalizedEmail:
    return NormalizedEmail(
        message_id="AAMkAGI2",
        mailbox="supply@ourcompany.example.com",
        sender=EmailAddress(address=sender),
        subject="VSL: NORTH STAR",
        body_text=BODY,
        attachments=[Attachment(filename="Requisition.xlsx")],
        **overrides,
    )


def outcome(
    category: EmailCategory = EmailCategory.NEW_RFQ,
    action: RecommendedAction = RecommendedAction.FORWARD_TO_DST,
    priority: Priority = Priority.NORMAL,
) -> ClassificationOutcome:
    return ClassificationOutcome(
        result=ClassificationResult(
            category=category,
            direction=Direction.INBOUND_CUSTOMER,
            requires_action=True,
            is_rfq=category is EmailCategory.NEW_RFQ,
            recommended_action=action,
            confidence=0.95,
            needs_human_review=False,
            priority=priority,
            decision_path=DecisionPath.LLM,
            reasoning="The customer asks the chandler to quote an attached requisition.",
            evidence=["'find attached our RFQ'"],
            extracted=Signals(vessel_name="NORTH STAR"),
        ),
        thread=SplitThread(latest_message=BODY),
        hints=Hints(),
        model="gpt-5.6-luna",
    )


async def one_rfq(tmp_path: Path, **kwargs) -> tuple[EmailRecords, str]:
    """A recorded email, labelled and forwarded, as a finished run leaves it."""
    records = EmailRecords(tmp_path / "Database", enabled=True)
    record_id = await records.open(
        email=email(), outcome=outcome(**kwargs), decision_id="be34a9b8-e2d3", source="outlook"
    )
    assert record_id is not None
    await records.update(
        record_id,
        labels=["SSG RFQ"],
        labelled=True,
        delivery=RecordedDelivery(outcome=DeliveryOutcome.SENT.value, forwarded_to="uae@desk"),
        files=[SourceFile(filename="Requisition.xlsx", data=b"rows", size_bytes=4)],
    )
    return records, record_id


async def test_one_row_per_email_with_the_label_and_the_sender(tmp_path: Path):
    records, record_id = await one_rfq(tmp_path)

    body = client(records).get(URL).json()

    assert body["total"] == 1
    row = body["items"][0]
    assert row["labels"] == ["SSG RFQ"]
    assert row["customerName"] == "purchasing@newcompany.example.com"
    assert row["rowKey"] == record_id


async def test_the_customer_s_own_words_never_reach_the_page(tmp_path: Path):
    """The record holds the body, the reasoning and the evidence. The row is
    what a browser gets, and it must carry none of them."""
    records, _ = await one_rfq(tmp_path)

    row = client(records).get(URL).json()["items"][0]

    printed = str(row)
    assert BODY not in printed
    assert "requisition" not in printed.lower()
    assert "subject" not in row and "reasoning" not in row


async def test_every_column_the_agent_cannot_answer_is_empty(tmp_path: Path):
    """Not zero, and not a plausible-looking guess: a table that reads as
    finished work is worse than one with blanks in it."""
    records, _ = await one_rfq(tmp_path)

    row = client(records).get(URL).json()["items"][0]

    assert row["counts"] is None
    assert row["completionPercent"] is None
    assert row["quotationNumber"] == ""
    assert row["responsibleUser"] == ""
    assert row["processingTime"] == ""
    assert row["receivedOn"] == ""
    # But it can be opened: every row on this list is an RFQ the agent read.
    assert row["id"] == row["rowKey"]


async def test_the_list_is_rfqs_and_not_the_rest_of_the_mailbox(tmp_path: Path):
    """The agent classifies everything that arrives; this page is about
    quoting. A supplier's reply and a marketing blast belong in the record
    folder and on the journal line, not here."""
    records = EmailRecords(tmp_path / "Database", enabled=True)
    forwarded = await records.open(
        email=email(), outcome=outcome(), decision_id="aaaa-1", source="outlook"
    )
    await records.update(
        forwarded,
        labels=["SSG RFQ"],
        delivery=RecordedDelivery(outcome=DeliveryOutcome.SENT.value),
    )
    await records.open(
        email=email("noreply@marketing.example.invalid"),
        outcome=outcome(EmailCategory.SPAM_MARKETING, RecommendedAction.IGNORE),
        decision_id="bbbb-2",
        source="outlook",
    )

    body = client(records).get(URL).json()

    assert body["total"] == 1
    assert body["items"][0]["customerName"] == "purchasing@newcompany.example.com"
    assert body["items"][0]["status"] == "inProgress"


async def test_a_row_counts_the_lines_its_matching_screen_will_show(tmp_path: Path):
    """The same number the page prints as "3 of 3 line(s) matched", so the list
    and the screen behind it never disagree about how big the RFQ is."""
    records, record_id = await one_rfq(tmp_path)
    await records.update(
        record_id,
        matching=[RecordedMatch(index=1, verbatim="bolts"), RecordedMatch(index=2, verbatim="rope")],
    )

    assert client(records).get(URL).json()["items"][0]["lineCount"] == 2


async def test_an_rfq_nobody_matched_counts_no_lines(tmp_path: Path):
    """Zero rather than null: this is an RFQ the agent read, and its matching
    screen really does have nothing on it."""
    records, _ = await one_rfq(tmp_path)

    assert client(records).get(URL).json()["items"][0]["lineCount"] == 0


async def test_an_urgent_email_carries_its_priority(tmp_path: Path):
    records, _ = await one_rfq(tmp_path, priority=Priority.URGENT)

    row = client(records).get(URL).json()["items"][0]

    assert row["priority"] == "HIGH"


async def test_search_matches_the_two_columns_that_have_anything_in_them(
    tmp_path: Path,
):
    records = EmailRecords(tmp_path / "Database", enabled=True)
    for index, sender in enumerate(["customer@almi.example.com", "noreply@portal.example.invalid"]):
        record_id = await records.open(
            email=email(sender), outcome=outcome(), decision_id=f"cccc-{index}", source="outlook"
        )
        await records.update(record_id, labels=["SSG RFQ" if index == 0 else "SSG No action"])

    found = client(records).get(URL, params={"search": "customer"}).json()

    assert found["total"] == 1
    assert found["items"][0]["customerName"] == "customer@almi.example.com"


async def test_pages_are_pages(tmp_path: Path):
    records = EmailRecords(tmp_path / "Database", enabled=True)
    for index in range(3):
        await records.open(
            email=email(f"customer{index}@sea.com"),
            outcome=outcome(),
            decision_id=f"dddd-{index}",
            source="outlook",
        )

    body = client(records).get(URL, params={"page": 2, "pageSize": 2}).json()

    assert body["total"] == 3
    assert len(body["items"]) == 1
    assert body["page"] == 2 and body["pageSize"] == 2


async def test_an_empty_database_is_an_empty_table_not_an_error(tmp_path: Path):
    records = EmailRecords(tmp_path / "Database", enabled=True)

    response = client(records).get(URL)

    assert response.status_code == 200
    assert response.json() == {"items": [], "total": 0, "page": 1, "pageSize": 25}


async def test_an_rfq_reads_as_its_lines(tmp_path: Path):
    """The screen compares two halves of each line - what the customer asked
    for, and what we sell - so the wire keeps them apart the way the page does."""
    records, record_id = await one_rfq(tmp_path)
    await records.update(
        record_id,
        matching=[
            RecordedMatch(
                index=1,
                verbatim="Hexagon Head Bolts (Bolt with Nut) M16*65",
                query="",
                customer_code="691284",
                quantity="500",
                uom="set",
                item_code="T69128400",
                item_description="HEX HEAD BOLT/NUT STEEL UNGALV, M16 X 65MM",
                confidence=98,
                item={"Item Code": "T69128400", "UOM": "SET", "Price": "0.42"},
                how="code_confirmed",
                why="Same bolt, same size.",
                candidates=[
                    RecordedCandidate(item_code="T69128400", description="HEX...", confidence=98),
                    RecordedCandidate(item_code="T69133100", description="HEX...", confidence=20),
                ],
            )
        ],
    )

    body = client(records).get(f"{URL}/{record_id}/rfq").json()

    assert body["customerName"] == "purchasing@newcompany.example.com"
    assert body["vesselName"] == ""  # this record carries no extraction
    line = body["lines"][0]
    assert line["line"] == 1
    assert (line["customerCode"], line["quantity"], line["uom"]) == ("691284", "500", "set")
    assert line["customerDescription"] == "Hexagon Head Bolts (Bolt with Nut) M16*65"
    assert line["itemCode"] == "T69128400"
    assert line["confidence"] == 98
    # Every candidate, scored - a refusal only means something beside what it
    # refused, and this is what the review panel reads.
    assert [one["confidence"] for one in line["candidates"]] == [98, 20]
    # The whole row of the sheet, as the sheet has it **now** - not the copy
    # the record took when it matched this line. The supplier and the price are
    # edited in the sheet, and a screen somebody settles a line from may not
    # show last month's copy of them.
    assert line["item"]["Price"] == "0.51", "0.42 is what the record kept"
    assert line["item"]["Supplier"] == "Northgate Marine Fasteners Ltd."


async def _rfq_with_a_shortlist(tmp_path: Path):
    """One RFQ, one line, one proposal and one alternative behind it."""
    records, record_id = await one_rfq(tmp_path)
    await records.update(
        record_id,
        matching=[
            RecordedMatch(
                index=7,
                verbatim="bolts hex head with nuts, full thread",
                customer_code="",
                how="search",
                candidates=[
                    RecordedCandidate(item_code="T69128400", confidence=100),
                    RecordedCandidate(item_code="T69133100", confidence=78),
                ],
            )
        ],
    )
    return records, record_id


def _confirmation(record_id: str, index: int) -> str:
    return f"{URL}/{record_id}/rfq/lines/{index}/confirmation"


async def test_a_line_is_settled_on_one_of_its_candidates(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)

    assert page.put(_confirmation(record_id, 7), json={"itemCode": "T69133100"}).status_code == 204

    line = page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]
    assert line["confirmedItemCode"] == "T69133100"


async def test_the_wire_carries_the_record_s_own_line_number(tmp_path: Path):
    """The page numbers rows 1..n; the record keeps the reader's numbering, and
    a confirmation is addressed by the second one. A screen that renumbers must
    not re-point a confirmation at somebody else's product."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)

    line = client(records).get(f"{URL}/{record_id}/rfq").json()["lines"][0]

    assert (line["line"], line["index"]) == (1, 7)


async def test_a_product_nobody_sells_is_a_404(tmp_path: Path):
    """The check is the sheet, not the shortlist: a code no row carries must
    never reach an order."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)

    assert page.put(_confirmation(record_id, 7), json={"itemCode": "T00000000"}).status_code == 404
    assert page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]["confirmedItemCode"] == ""


async def test_a_product_picked_by_hand_is_settled_though_it_was_never_offered(tmp_path: Path):
    """The five candidates are a proposal. Somebody who finds none of them
    right goes and picks the sixth, and that is the point of the picker."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)

    assert page.put(_confirmation(record_id, 7), json={"itemCode": "T85116300"}).status_code == 204

    line = page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]
    assert line["confirmedItemCode"] == "T85116300"
    assert [one["itemCode"] for one in line["candidates"]] == ["T69128400", "T69133100"]


async def test_changing_your_mind_clears_the_confirmation(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)
    page.put(_confirmation(record_id, 7), json={"itemCode": "T69128400"})

    assert page.delete(_confirmation(record_id, 7)).status_code == 204
    assert page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]["confirmedItemCode"] == ""


async def test_a_line_nobody_has_cannot_be_settled(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)

    response = client(records).put(_confirmation(record_id, 99), json={"itemCode": "T69128400"})

    assert response.status_code == 404


async def test_an_rfq_nobody_has_is_a_404(tmp_path: Path):
    records, _ = await one_rfq(tmp_path)

    assert client(records).get(f"{URL}/nothing-here/rfq").status_code == 404


async def test_one_email_can_be_read_in_full(tmp_path: Path):
    """The row is deliberately thin; this is where the whole record lives."""
    records, record_id = await one_rfq(tmp_path)

    body = client(records).get(f"{URL}/{record_id}").json()

    assert body["subject"] == "VSL: NORTH STAR"
    assert body["verdict"]["category"] == "NEW_RFQ"
    # The record goes out as it is on disk - this is the file itself, not a
    # view of it, and the folder is meant to be readable by a person.
    assert body["attachments"][0]["saved_as"] == "Requisition.xlsx"


async def test_an_attachment_comes_back_as_the_bytes_that_arrived(tmp_path: Path):
    records, record_id = await one_rfq(tmp_path)

    response = client(records).get(f"{URL}/{record_id}/files/Requisition.xlsx")

    assert response.status_code == 200
    assert response.content == b"rows"


async def test_the_stream_says_when_the_list_changed():
    """The page is told to read the list again; it is never told what to think.

    Driven through the generator rather than over HTTP: the response never
    ends, and a test client reading it to the end waits forever.
    """
    changes = Changes()
    stream = quotes.events(changes)
    waiting = asyncio.ensure_future(anext(stream))
    await asyncio.sleep(0)  # let it subscribe before anything is announced

    changes.announce()

    assert await asyncio.wait_for(waiting, timeout=1) == quotes.CHANGED
    await stream.aclose()


async def test_a_silent_stream_still_says_something(monkeypatch):
    """A connection that sends nothing for minutes is closed by every proxy
    between the page and the agent."""
    monkeypatch.setattr(quotes, "KEEPALIVE_SECONDS", 0.01)
    stream = quotes.events(Changes())

    assert await asyncio.wait_for(anext(stream), timeout=1) == quotes.KEEPALIVE
    await stream.aclose()


async def test_a_page_that_goes_away_stops_being_a_listener():
    """A closed browser tab must not leave an event behind to be set forever."""
    changes = Changes()
    stream = quotes.events(changes)
    waiting = asyncio.ensure_future(anext(stream))
    await asyncio.sleep(0)  # let it subscribe
    assert changes.listeners == 1

    # What a disconnected client does to the response body.
    waiting.cancel()
    with suppress(asyncio.CancelledError):
        await waiting

    assert changes.listeners == 0


async def test_a_path_out_of_the_record_is_a_404(tmp_path: Path):
    records, record_id = await one_rfq(tmp_path)

    assert client(records).get(f"{URL}/{record_id}/files/../../../.env").status_code == 404
    assert client(records).get(f"{URL}/nothing-here").status_code == 404


async def test_a_settled_line_is_shown_as_what_it_was_settled_on(tmp_path: Path):
    """The record keeps only the code. Its description, source and unit come
    from today's sheet - a page showing last week's copy of them would be
    showing something nobody sells."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)
    page.put(_confirmation(record_id, 7), json={"itemCode": "T85116300"})

    line = page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]

    assert line["itemCode"] == "T85116300"
    assert line["itemDescription"] == "WELDER GLOVES FIVE FINGERS"
    assert line["confidence"] is None, "a person chose it; the number was not the reason"


async def test_a_settled_candidate_keeps_the_score_it_was_shortlisted_with(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)
    page.put(_confirmation(record_id, 7), json={"itemCode": "T69133100"})

    line = page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]

    assert (line["itemCode"], line["confidence"]) == ("T69133100", 78)


async def test_every_candidate_says_why_it_scored_what_it_did(tmp_path: Path):
    records, record_id = await one_rfq(tmp_path)
    await records.update(
        record_id,
        matching=[
            RecordedMatch(
                index=7,
                verbatim="Steel toe sneakers",
                how="search",
                candidates=[
                    RecordedCandidate(item_code="T69128400", confidence=63, why="size not stated"),
                    RecordedCandidate(item_code="T69133100", confidence=None),
                ],
            )
        ],
    )

    candidates = client(records).get(f"{URL}/{record_id}/rfq").json()["lines"][0]["candidates"]

    assert [(one["confidence"], one["why"]) for one in candidates] == [
        (63, "size not stated"),
        (None, ""),
    ]


async def test_every_candidate_carries_the_sheet_s_row_as_it_stands_today(tmp_path: Path):
    """A candidate is a product somebody is about to choose between, so it is
    shown as the sheet has it now - which is where the supplier lives."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)

    line = client(records).get(f"{URL}/{record_id}/rfq").json()["lines"][0]

    assert [one["item"]["Supplier"] for one in line["candidates"]] == [
        "Northgate Marine Fasteners Ltd.",
        "Northgate Marine Fasteners Ltd.",
    ]


async def test_a_product_the_sheet_has_dropped_keeps_the_row_the_record_kept(tmp_path: Path):
    """Fewer columns beats a candidate that goes blank: a code the sheet no
    longer carries is shown as the record last saw it."""
    records, record_id = await one_rfq(tmp_path)
    await records.update(
        record_id,
        matching=[
            RecordedMatch(
                index=1,
                verbatim="turbocharger cartridge NR34/S",
                how="search",
                candidates=[
                    RecordedCandidate(
                        item_code="T00000000",
                        confidence=40,
                        item={"Item Code": "T00000000", "Supplier": "A firm we no longer list"},
                    )
                ],
            )
        ],
    )

    line = client(records).get(f"{URL}/{record_id}/rfq").json()["lines"][0]

    assert line["candidates"][0]["item"]["Supplier"] == "A firm we no longer list"


def _offer(record_id: str, index: int) -> str:
    return f"{URL}/{record_id}/rfq/lines/{index}/offer"


async def test_a_line_carries_what_a_supplier_quoted_for_one_unit(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)

    assert page.put(_offer(record_id, 7), json={"unitPrice": 24.5}).status_code == 204

    line = page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]
    assert line["offerUnitPrice"] == 24.5


async def test_a_line_nobody_quoted_carries_no_price_rather_than_zero(tmp_path: Path):
    """Zero is a price somebody named. Null is nobody having answered, and the
    pricing screen may not read the second as the first."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)

    line = client(records).get(f"{URL}/{record_id}/rfq").json()["lines"][0]

    assert line["offerUnitPrice"] is None


async def test_a_withdrawn_offer_leaves_the_line_unpriced(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)
    page.put(_offer(record_id, 7), json={"unitPrice": 24.5})

    assert page.delete(_offer(record_id, 7)).status_code == 204
    assert page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]["offerUnitPrice"] is None


async def test_a_price_no_supplier_could_have_named_is_refused(tmp_path: Path):
    """Both ends of it. A zero would read as "free" on the pricing screen, and
    a typo with four extra digits as a quotation somebody has to explain."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)

    assert page.put(_offer(record_id, 7), json={"unitPrice": 0}).status_code == 422
    assert page.put(_offer(record_id, 7), json={"unitPrice": -1}).status_code == 422
    assert page.put(_offer(record_id, 7), json={"unitPrice": 1e12}).status_code == 422
    assert page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]["offerUnitPrice"] is None


async def test_a_line_nobody_has_cannot_be_priced(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)

    assert client(records).put(_offer(record_id, 99), json={"unitPrice": 24.5}).status_code == 404


async def test_a_price_and_a_confirmation_do_not_disturb_each_other(tmp_path: Path):
    """Two separate decisions on one line: which product, and what it costs."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)

    page.put(_confirmation(record_id, 7), json={"itemCode": "T69133100"})
    page.put(_offer(record_id, 7), json={"unitPrice": 24.5})

    line = page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]
    assert (line["confirmedItemCode"], line["offerUnitPrice"]) == ("T69133100", 24.5)


async def test_an_unsettled_line_still_shows_what_the_agent_found(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)

    line = client(records).get(f"{URL}/{record_id}/rfq").json()["lines"][0]

    assert line["confirmedItemCode"] == ""
    assert line["itemCode"] == "", "this line was searched, so the agent settled nothing"


async def test_a_priced_line_carries_the_moment_the_price_arrived(tmp_path: Path):
    """The responses screen dates a supplier's reply by this."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)

    page.put(_offer(record_id, 7), json={"unitPrice": 24.5})

    assert page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]["offerReceivedAt"] is not None


async def test_an_unpriced_line_carries_no_moment(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)

    line = client(records).get(f"{URL}/{record_id}/rfq").json()["lines"][0]
    assert line["offerReceivedAt"] is None


async def test_a_withdrawn_offer_takes_its_moment_with_it(tmp_path: Path):
    """A line with no price but a date on it would read as a reply we lost."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)
    page.put(_offer(record_id, 7), json={"unitPrice": 24.5})

    page.delete(_offer(record_id, 7))

    assert page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]["offerReceivedAt"] is None


def _inquiries(record_id: str) -> str:
    return f"{URL}/{record_id}/rfq/inquiries"


ASKED = {
    "inquiries": [
        {
            "supplier": "Marinet Services Pte Ltd",
            "body": "Dear Marinet Services Pte Ltd,\n\nKindly quote the following item(s).",
            "lines": [7],
        }
    ]
}


async def test_an_rfq_starts_with_nobody_having_been_asked(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)

    assert client(records).get(f"{URL}/{record_id}/rfq").json()["inquiries"] == []


async def test_the_letter_that_went_to_a_supplier_is_kept_whole(tmp_path: Path):
    """The text is editable before it goes, so the sentence that was sent is
    the one worth keeping - not the one the template would rebuild."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)

    assert page.put(_inquiries(record_id), json=ASKED).status_code == 204

    sent = page.get(f"{URL}/{record_id}/rfq").json()["inquiries"]
    assert len(sent) == 1
    assert sent[0]["supplier"] == "Marinet Services Pte Ltd"
    assert sent[0]["body"].startswith("Dear Marinet Services Pte Ltd,")
    assert sent[0]["lines"] == [7]
    assert sent[0]["sentAt"] is not None, "the server dates it, not the browser"


async def test_the_inquiries_go_out_once(tmp_path: Path):
    """The screen greys its button out for the same reason. A rule that lives
    only in a button is not one."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)
    page.put(_inquiries(record_id), json=ASKED)

    second = {"inquiries": [{"supplier": "Hansa Technik", "body": "Dear Hansa,", "lines": [7]}]}
    assert page.put(_inquiries(record_id), json=second).status_code == 409

    sent = page.get(f"{URL}/{record_id}/rfq").json()["inquiries"]
    assert [one["supplier"] for one in sent] == ["Marinet Services Pte Ltd"]


async def test_every_letter_of_one_send_lands_together(tmp_path: Path):
    """One click sends them all, and half of them on the record would describe
    a send that never happened."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)

    page.put(
        _inquiries(record_id),
        json={
            "inquiries": [
                {"supplier": "Marinet Services Pte Ltd", "body": "Dear Marinet,", "lines": [7]},
                {"supplier": "Hansa Technik", "body": "Dear Hansa,", "lines": [7]},
            ]
        },
    )

    sent = page.get(f"{URL}/{record_id}/rfq").json()["inquiries"]
    assert {one["supplier"] for one in sent} == {"Marinet Services Pte Ltd", "Hansa Technik"}


async def test_an_rfq_that_is_not_there_cannot_be_asked(tmp_path: Path):
    records, _ = await _rfq_with_a_shortlist(tmp_path)

    assert client(records).put(_inquiries("no-such-rfq"), json=ASKED).status_code == 409


async def test_a_letter_that_could_not_be_one_is_refused(tmp_path: Path):
    """Guards rather than rules: what a person writes is theirs, and these
    only refuse what no letter could be."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)

    assert page.put(_inquiries(record_id), json={"inquiries": []}).status_code == 422
    assert (
        page.put(
            _inquiries(record_id), json={"inquiries": [{"supplier": "", "body": "x", "lines": []}]}
        ).status_code
        == 422
    )
    assert (
        page.put(
            _inquiries(record_id), json={"inquiries": [{"supplier": "A", "body": "", "lines": []}]}
        ).status_code
        == 422
    )
    assert page.get(f"{URL}/{record_id}/rfq").json()["inquiries"] == []


async def test_asking_the_suppliers_leaves_the_lines_alone(tmp_path: Path):
    """Two different screens, and neither write may disturb the other."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)
    page.put(_confirmation(record_id, 7), json={"itemCode": "T69133100"})

    page.put(_inquiries(record_id), json=ASKED)

    assert page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]["confirmedItemCode"] == "T69133100"


def _approval(record_id: str) -> str:
    return f"{URL}/{record_id}/rfq/approval"


SIGNED = {"marginStock": 12, "marginJit": 15, "lines": [{"index": 7, "unitPrice": 28.0}]}


async def test_an_rfq_starts_unapproved(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)

    page = client(records).get(f"{URL}/{record_id}/rfq").json()
    assert page["approval"] is None
    assert page["lines"][0]["approvedUnitPrice"] is None


async def test_approving_freezes_the_prices_and_the_margins(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)

    assert page.put(_approval(record_id), json=SIGNED).status_code == 204

    rfq = page.get(f"{URL}/{record_id}/rfq").json()
    assert rfq["approval"]["marginStock"] == 12
    assert rfq["approval"]["marginJit"] == 15
    assert rfq["approval"]["approvedAt"] is not None, "the server dates it, not the browser"
    assert rfq["lines"][0]["approvedUnitPrice"] == 28.0


async def test_an_approval_with_a_hole_in_it_is_refused(tmp_path: Path):
    """A quotation missing a line is not a smaller quotation - it is a wrong
    one, and the hole is invisible in the total."""
    records, record_id = await one_rfq(tmp_path)
    await records.update(
        record_id,
        matching=[RecordedMatch(index=7, verbatim="bolts"), RecordedMatch(index=8, verbatim="rope")],
    )
    page = client(records)

    assert page.put(_approval(record_id), json=SIGNED).status_code == 422
    assert page.get(f"{URL}/{record_id}/rfq").json()["approval"] is None


async def test_the_pricing_is_approved_once(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)
    page.put(_approval(record_id), json=SIGNED)

    again = {"marginStock": 99, "marginJit": 99, "lines": [{"index": 7, "unitPrice": 99.0}]}
    assert page.put(_approval(record_id), json=again).status_code == 409

    rfq = page.get(f"{URL}/{record_id}/rfq").json()
    assert rfq["approval"]["marginStock"] == 12
    assert rfq["lines"][0]["approvedUnitPrice"] == 28.0


async def test_an_rfq_that_is_not_there_cannot_be_approved(tmp_path: Path):
    records, _ = await _rfq_with_a_shortlist(tmp_path)

    assert client(records).put(_approval("no-such-rfq"), json=SIGNED).status_code == 404


async def test_a_price_or_a_margin_no_desk_could_have_meant_is_refused(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)

    assert page.put(_approval(record_id), json={**SIGNED, "marginStock": -1}).status_code == 422
    assert page.put(_approval(record_id), json={**SIGNED, "marginJit": 99999}).status_code == 422
    assert page.put(_approval(record_id), json={**SIGNED, "lines": []}).status_code == 422
    assert page.get(f"{URL}/{record_id}/rfq").json()["approval"] is None


async def test_an_approved_price_does_not_follow_the_supplier_afterwards(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)
    page.put(_approval(record_id), json=SIGNED)

    page.put(_offer(record_id, 7), json={"unitPrice": 99.0})

    line = page.get(f"{URL}/{record_id}/rfq").json()["lines"][0]
    assert (line["offerUnitPrice"], line["approvedUnitPrice"]) == (99.0, 28.0)


def _quotation(record_id: str) -> str:
    return f"{URL}/{record_id}/rfq/quotation.pdf"


def _pdf_text(content: bytes) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)


async def _approved(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)
    page = client(records)
    page.put(_confirmation(record_id, 7), json={"itemCode": "T69133100"})
    page.put(
        _approval(record_id),
        json={"marginStock": 12, "marginJit": 15, "lines": [{"index": 7, "unitPrice": 28.0}]},
    )
    return page, record_id


async def test_the_quotation_downloads_as_a_pdf(tmp_path: Path):
    page, record_id = await _approved(tmp_path)

    response = page.get(_quotation(record_id))

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith("attachment; filename=")
    assert response.headers["content-disposition"].endswith('_quotation.pdf"')
    assert response.content.startswith(b"%PDF")


async def test_the_quotation_quotes_the_approved_price(tmp_path: Path):
    """The document is the number we named, not one worked out again today."""
    page, record_id = await _approved(tmp_path)
    page.put(_offer(record_id, 7), json={"unitPrice": 99.0})

    text = _pdf_text(page.get(_quotation(record_id)).content)

    assert "T69133100" in text, "our code, as settled"
    assert "28.00" in text
    assert "99.00" not in text


async def test_there_is_no_quotation_before_the_pricing_is_approved(tmp_path: Path):
    """The screen keeps the fourth stage shut until approval. A rule that
    lives only in a button is not one."""
    records, record_id = await _rfq_with_a_shortlist(tmp_path)

    assert client(records).get(_quotation(record_id)).status_code == 409


async def test_there_is_no_quotation_for_an_rfq_that_is_not_there(tmp_path: Path):
    records, _ = await _rfq_with_a_shortlist(tmp_path)

    assert client(records).get(_quotation("no-such-rfq")).status_code == 404


async def test_the_uae_letterhead_goes_out_from_the_dubai_office(tmp_path: Path):
    page, record_id = await _approved(tmp_path)

    response = page.get(_quotation(record_id), params={"format": "uae"})

    assert response.status_code == 200
    text = _pdf_text(response.content)
    assert "Seven Seas Shipchandlers (L.L.C)" in text
    assert "VAT%" in text
    assert "28.00" in text


async def test_a_letterhead_nobody_has_is_refused(tmp_path: Path):
    page, record_id = await _approved(tmp_path)

    assert page.get(_quotation(record_id), params={"format": "fr"}).status_code == 422


def _customer_file(record_id: str) -> str:
    return f"{URL}/{record_id}/rfq/customer-file.xlsx"


async def test_the_customer_file_downloads_as_a_spreadsheet(tmp_path: Path):
    page, record_id = await _approved(tmp_path)

    response = page.get(_customer_file(record_id))

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"].endswith('_customer_file.xlsx"')


async def test_the_customer_file_carries_the_approved_price(tmp_path: Path):
    """The number we named, not one worked out again after a supplier moved."""
    page, record_id = await _approved(tmp_path)
    page.put(_offer(record_id, 7), json={"unitPrice": 99.0})

    sheet = load_workbook(BytesIO(page.get(_customer_file(record_id)).content)).active
    row = next(sheet.iter_rows(min_row=2, values_only=True))

    assert row[0] is None, "no category - we have none to give"
    assert row[1] == 1
    assert row[6] == 28.0


async def test_there_is_no_customer_file_before_the_pricing_is_approved(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)

    assert client(records).get(_customer_file(record_id)).status_code == 409


async def test_there_is_no_customer_file_for_an_rfq_that_is_not_there(tmp_path: Path):
    records, _ = await _rfq_with_a_shortlist(tmp_path)

    assert client(records).get(_customer_file("no-such-rfq")).status_code == 404


def _quote_workbook(record_id: str) -> str:
    return f"{URL}/{record_id}/rfq/quotation.xlsm"


async def test_the_quote_workbook_downloads_for_either_office(tmp_path: Path):
    page, record_id = await _approved(tmp_path)

    for office in ("sg", "uae"):
        response = page.get(_quote_workbook(record_id), params={"format": office})

        assert response.status_code == 200
        assert response.headers["content-type"] == "application/vnd.ms-excel.sheet.macroEnabled.12"
        assert response.headers["content-disposition"].endswith(f'_quotation_{office}.xlsm"')


async def test_the_quote_workbook_carries_the_approved_price_and_the_office(tmp_path: Path):
    page, record_id = await _approved(tmp_path)
    page.put(_offer(record_id, 7), json={"unitPrice": 99.0})

    content = page.get(_quote_workbook(record_id), params={"format": "uae"}).content
    workbook = load_workbook(BytesIO(content), keep_vba=True)
    lines = workbook[workbook.sheetnames[1]]

    assert workbook["SUMMARY"]["C15"].value == "100569476300003"
    assert lines["C11"].value == "T69133100"
    assert lines["K11"].value == 28


async def test_there_is_no_quote_workbook_before_the_pricing_is_approved(tmp_path: Path):
    records, record_id = await _rfq_with_a_shortlist(tmp_path)

    assert client(records).get(_quote_workbook(record_id)).status_code == 409


async def test_there_is_no_quote_workbook_for_an_office_we_do_not_have(tmp_path: Path):
    page, record_id = await _approved(tmp_path)

    assert page.get(_quote_workbook(record_id), params={"format": "fr"}).status_code == 422

