"""The quotation in the customer's own spreadsheet layout.

Read back with openpyxl. What is held down is what the customer will act on:
the header row they know, one row per line in their order, their codes intact,
and the approved numbers as numbers.
"""

from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from src.services.customer_file import fill
from src.services.quotation import QuotedLine
from src.tools.make_customer_template import HEADINGS, build

TEMPLATE = Path(__file__).parents[1] / "config" / "quotation_customer_file_template.xlsx"


def line(number: int = 1, **overrides) -> QuotedLine:
    fields = {
        "number": number,
        "code": "04361753",
        "description": "CHIN STRAP FOR MUNDO SAFETY HELMET (SF-06B)",
        "quantity": "10",
        "uom": "pcs",
        "unit_price": Decimal("28.00"),
    }
    return QuotedLine(**{**fields, **overrides})


def sheet_of(workbook: bytes):
    return load_workbook(BytesIO(workbook)).active


def rows_of(workbook: bytes) -> list[tuple]:
    return list(sheet_of(workbook).iter_rows(values_only=True))


def test_the_shipped_template_is_what_the_tool_builds(tmp_path: Path):
    """The file in `config/` is the tool's output, not a hand-edited copy that
    drifted from it."""
    fresh = tmp_path / "template.xlsx"
    build(fresh)

    assert rows_of(fresh.read_bytes()) == rows_of(TEMPLATE.read_bytes())


def test_it_keeps_the_customer_s_header_row():
    assert rows_of(fill(TEMPLATE.read_bytes(), [line()]))[0] == HEADINGS


def test_one_row_per_line_in_order():
    rows = rows_of(fill(TEMPLATE.read_bytes(), [line(1), line(2, code="851163")]))

    assert [row[1] for row in rows[1:]] == [1, 2]
    assert rows[2][2] == "851163"


def test_the_category_is_left_blank():
    """We have no category for a line, and a guessed one would be ours, not
    the customer's."""
    assert rows_of(fill(TEMPLATE.read_bytes(), [line()]))[1][0] is None


def test_a_code_keeps_its_leading_zero():
    assert rows_of(fill(TEMPLATE.read_bytes(), [line()]))[1][2] == "04361753"


def test_the_numbers_are_numbers():
    """So the customer can sum or sort the column without retyping it."""
    priced = line(quantity="100", unit_price=Decimal("48.65"))
    row = rows_of(fill(TEMPLATE.read_bytes(), [priced]))[1]

    assert row[5:] == (100, 48.65, 4865)


def test_a_quantity_that_is_not_a_number_is_kept_as_written_and_not_totalled():
    row = rows_of(fill(TEMPLATE.read_bytes(), [line(quantity="2 coil")]))[1]

    assert row[5] == "2 coil"
    assert row[7] is None


def test_every_row_takes_the_template_s_look():
    sheet = sheet_of(fill(TEMPLATE.read_bytes(), [line(1), line(2), line(3)]))

    for row in (2, 3, 4):
        assert sheet.cell(row=row, column=4).fill.fgColor.rgb == "FFE0F9F9"
        assert sheet.cell(row=row, column=3).number_format == "@"


def test_no_lines_leaves_no_empty_row():
    """A styled empty row reads as a line somebody forgot to fill in."""
    assert rows_of(fill(TEMPLATE.read_bytes(), [])) == [HEADINGS]
