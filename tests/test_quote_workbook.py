"""The desk's own quotation workbook, `Quote.xlsm`, filled.

Read back with openpyxl for the values and as a zip for everything openpyxl
would not keep - the macro, its button, the printer settings. What is held
down: the office block each letterhead fixes, the quotation's own data, one
live row per line, and nothing of the customer the template was made from.
"""

import io
import re
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.services.quote_workbook import (
    DUBAI_OFFICE,
    SINGAPORE_OFFICE,
    BookLine,
    QuoteBook,
    fill,
    sheet_name,
)
from src.tools.make_quote_template import build, restyle, uniform_fonts

CONFIG = Path(__file__).parents[1] / "config"
TEMPLATE = CONFIG / "quotation_sg_uae_workbook_template.xlsm"
# The real export the blank was made from. Kept out of git (it is a real
# customer's), so the check that rebuilds the blank runs only where it is.
EXPORT = Path(__file__).parents[1] / "docs" / "samples" / "Quote.xlsm"

# The parts openpyxl drops, and the reason the workbook is edited as a zip.
KEPT = (
    "xl/vbaProject.bin",
    "xl/ctrlProps/ctrlProp1.xml",
    "xl/drawings/GemVmlDrawing4102.vml",
    "xl/printerSettings/printerSettings1.bin",
    "xl/media/GemImage2438.jpeg",
)


def line(number: int = 1, **overrides) -> BookLine:
    fields = {
        "number": number,
        "code": "T33116401",
        "description": "CHIN STRAP FOR SAFETY HELMETS",
        "quantity": "10",
        "uom": "pcs",
        "unit_price": Decimal("28.00"),
        "customer_code": "04361753",
        "customer_description": "CHIN STRAP FOR MUNDO SAFETY HELMET (SF-06B)",
    }
    return BookLine(**{**fields, **overrides})


def book(*lines: BookLine, **overrides) -> QuoteBook:
    fields = {
        "office": DUBAI_OFFICE,
        "number": "RFQ-0002",
        "vessel": "MT ODENSE",
        "port": "Singapore",
        "client": "purchasing@example.com",
        "quoted_on": date(2026, 9, 25),
        "lines": lines or (line(),),
    }
    return QuoteBook(**{**fields, **overrides})


def filled(doc: QuoteBook) -> bytes:
    return fill(TEMPLATE.read_bytes(), doc)


def opened(workbook: bytes, *, cached: bool = False):
    return load_workbook(io.BytesIO(workbook), keep_vba=True, data_only=cached)


def test_the_template_carries_nothing_of_the_customer_it_was_made_from():
    """Made from a real export, for a real customer; none of it may travel."""
    with zipfile.ZipFile(TEMPLATE) as template:
        everything = b"".join(template.read(name) for name in template.namelist())

    # `C:\\Workspace` is the desk's disk path the export recorded; a bare
    # `[Workspace]` is a section of every VBA project and says nothing.
    for leftover in (b"OLDENDORFF", b"TEOL-260139-1", b"Sankara", b"BRIO", b"C:\\Workspace"):
        assert leftover not in everything


@pytest.mark.skipif(not EXPORT.is_file(), reason="the SCINT export is not in this checkout")
def test_the_shipped_template_is_what_the_tool_builds(tmp_path: Path):
    fresh = tmp_path / "template.xlsm"
    build(EXPORT, fresh)

    sheets = ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml", "xl/workbook.xml")
    with zipfile.ZipFile(fresh) as made, zipfile.ZipFile(TEMPLATE) as shipped:
        for part in sheets:
            assert made.read(part) == shipped.read(part), part


def test_the_macro_its_button_and_the_printer_settings_survive():
    with zipfile.ZipFile(io.BytesIO(filled(book()))) as copy:
        names = set(copy.namelist())
        with zipfile.ZipFile(TEMPLATE) as template:
            for part in KEPT:
                assert part in names
                assert copy.read(part) == template.read(part), part


def test_the_dubai_office_block_is_fixed():
    summary = opened(filled(book()))["SUMMARY"]

    assert summary["C9"].value == "Seven Seas Shipchandlers (L.L.C) (Dubai)"
    assert summary["C10"].value.startswith("Plot 598-668, Dubai Investments Park")
    assert summary["C13"].value == "Ph: +971 4 8033 3333"
    assert summary["C14"].value == "supply.uae@sevenseasgroup.com"
    assert summary["C15"].value == "100569476300003"
    assert summary["C20"].value == "60 DAYS FROM THE DATE OF INVOICE"


def test_the_singapore_office_block_is_fixed():
    summary = opened(filled(book(office=SINGAPORE_OFFICE)))["SUMMARY"]

    assert summary["C10"].value == (
        "12 Tuas Road, Singapore. 638486 UEN : 199305221C, Tax Reg No : M201169402"
    )
    assert summary["C13"].value == "Phone: +65 3152 2188 Fax: +65 3152 2189"
    assert summary["C14"].value == (
        "E-Mail: supply.singapore@sevenseasgroup.com INTERNET: www.sevenseasgroup.com"
    )
    assert summary["C15"].value == "M201169402"
    assert summary["C20"].value is None, "no office terms we could name"


def test_the_quotation_s_own_data_goes_in():
    summary = opened(filled(book()))["SUMMARY"]

    assert summary["C4"].value == "MT ODENSE"
    assert summary["C5"].value == "Singapore"
    assert summary["C7"].value == "RFQ-0002"
    assert summary["C8"].value == "25 September 2026"
    assert summary["C17"].value == "purchasing@example.com"
    assert summary["C19"].value is None, "no customer tax number we hold"


def test_the_line_sheet_is_named_after_the_quotation_and_the_summary_follows():
    workbook = opened(filled(book()))

    assert workbook.sheetnames == ["SUMMARY", "RFQ-0002"]
    assert workbook["SUMMARY"]["J7"].value == "='RFQ-0002'!L8"
    assert workbook["RFQ-0002"]["B1"].value == "RFQ-0002"
    assert workbook["RFQ-0002"]["D3"].value is None, "the customer's number is not ours to invent"


def test_one_live_row_per_line():
    """Formulas, not numbers: the desk changes a quantity and the book follows."""
    sheet = opened(filled(book(line(1), line(2, quantity="5"), line(3))))["RFQ-0002"]

    assert [sheet.cell(row=row, column=1).value for row in (11, 12, 13)] == [1, 2, 3]
    assert sheet["N12"].value == "=K12*I12"
    assert sheet["L13"].value == "=N13"
    assert sheet["L3"].value == "=SUM(N11:N13)"
    assert sheet["L6"].value == "=SUM(S11:S13)"
    assert sheet["A14"].value is None


def test_each_line_carries_both_sides():
    sheet = opened(filled(book()))["RFQ-0002"]
    row = [cell.value for cell in sheet[11]][:11]

    assert row == [
        1,
        "04361753",
        "04361753",
        "CHIN STRAP FOR MUNDO SAFETY HELMET (SF-06B)",
        None,
        None,
        "CHIN STRAP FOR SAFETY HELMETS",
        None,
        10,
        "pcs",
        28,
    ]


def test_every_row_keeps_the_model_row_s_look():
    sheet = opened(filled(book(line(1), line(2))))["RFQ-0002"]

    for column in ("A", "D", "K", "L"):
        assert sheet[f"{column}12"].style_id == sheet[f"{column}11"].style_id


def test_the_cache_shows_the_totals_before_excel_recalculates():
    """A previewer prints the cache, and a cache of zeros is a quotation for nothing."""
    doc = book(line(1), line(2, quantity="100", unit_price=Decimal("48.65")))
    workbook = opened(filled(doc), cached=True)

    assert workbook["RFQ-0002"]["L3"].value == 5145
    assert workbook["RFQ-0002"]["N12"].value == 4865
    assert workbook["SUMMARY"]["J18"].value == 5145
    assert workbook["SUMMARY"]["J19"].value == 0
    assert workbook["SUMMARY"]["J23"].value == 5145
    assert workbook["SUMMARY"]["H4"].value == "USD"


def test_excel_is_told_to_recalculate_on_open():
    with zipfile.ZipFile(io.BytesIO(filled(book()))) as copy:
        assert b'fullCalcOnLoad="1"' in copy.read("xl/workbook.xml")


def test_what_a_customer_wrote_is_written_as_text():
    sheet = opened(filled(book(line(customer_description='BOLTS <M12> & "NUTS"'))))["RFQ-0002"]

    assert sheet["D11"].value == 'BOLTS <M12> & "NUTS"'


def test_a_quantity_that_is_not_a_number_is_kept_as_written():
    sheet = opened(filled(book(line(quantity="2 coil"))), cached=True)["RFQ-0002"]

    assert sheet["I11"].value == "2 coil"


def test_a_sheet_name_is_one_excel_accepts():
    assert sheet_name("RFQ-0002") == "RFQ-0002"
    assert sheet_name("RFQ/00:02") == "RFQ0002"
    assert sheet_name("") == "QUOTE"
    assert sheet_name("summary") == "QUOTE", "the other sheet already has that name"
    assert len(sheet_name("R" * 40)) == 31


def test_every_formula_is_one_excel_itself_would_write():
    """The export stored `==H4` and `SUM(J20: J22)`. Desktop Excel repairs
    them; Excel for the web refuses the whole file."""
    with zipfile.ZipFile(io.BytesIO(filled(book()))) as copy:
        for part in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
            formulas = re.findall(r"<f[^>]*>(.*?)</f>", copy.read(part).decode("utf-8"))
            assert formulas, part
            for formula in formulas:
                assert not formula.startswith("="), formula
                assert ": " not in formula and " :" not in formula, formula


def _fonts(workbook: bytes) -> str:
    styles = zipfile.ZipFile(io.BytesIO(workbook)).read("xl/styles.xml").decode("utf-8")
    return re.search(r"<fonts\b.*?</fonts>", styles, re.DOTALL).group(0)


def test_one_font_and_one_size_in_the_whole_workbook():
    """The export mixed Arial 8 to 14 with Calibri 11 and 12; a quotation whose
    cells change size from one to the next reads as a mistake."""
    for office in (DUBAI_OFFICE, SINGAPORE_OFFICE):
        table = _fonts(filled(book(line(1), line(2), office=office)))

        assert set(re.findall(r'<name val="([^"]+)"', table)) == {"Arial"}
        assert set(re.findall(r'<sz val="([^"]+)"', table)) == {"10"}
        assert "scheme" not in table, "a theme font would let Excel put Calibri back"


def test_every_filled_cell_reads_arial_10():
    workbook = opened(filled(book(line(1), line(2))))

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    assert (cell.font.name, cell.font.sz) == ("Arial", 10), cell.coordinate


def test_uniform_fonts_keeps_bold_underline_and_colour():
    styles = (
        '<styleSheet><fonts count="2">'
        '<font><b /><u /><sz val="14" /><color rgb="FFFF0000" /><name val="Calibri" />'
        '<family val="2" /><scheme val="minor" /></font>'
        '<font><sz val="8" /><name val="Arial" /></font>'
        "</fonts><dxfs><dxf><font><color theme=\"0\" /></font></dxf></dxfs></styleSheet>"
    )

    fonts, _, rest = uniform_fonts(styles).partition("</fonts>")

    assert fonts.count('<sz val="10" /><name val="Arial" />') == 2
    assert '<b /><u /><color rgb="FFFF0000" />' in fonts
    assert "Calibri" not in fonts and "scheme" not in fonts and "family" not in fonts
    assert '<font><color theme="0" /></font>' in rest, "conditional formats are left alone"


def test_restyle_changes_nothing_but_the_fonts():
    before = TEMPLATE.read_bytes()
    after = restyle(before)

    with zipfile.ZipFile(io.BytesIO(before)) as old, zipfile.ZipFile(io.BytesIO(after)) as new:
        assert old.namelist() == new.namelist()
        for name in old.namelist():
            if name != "xl/styles.xml":
                assert old.read(name) == new.read(name), name

