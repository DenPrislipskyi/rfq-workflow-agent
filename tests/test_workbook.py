"""Stage C at the form level: which value belongs in which cell.

Nothing here decides anything - every value was read, verified and mapped
before it arrived. What is under test is the map itself, and the two rules the
task attaches to it: a cell this stage owns is written even when it is empty,
and the items go down in the order the customer wrote them.
"""

import io
import zipfile
from datetime import UTC, date, datetime
from pathlib import Path

import openpyxl
import pytest

from src.infrastructure.documents.models import (
    ATTACHMENT_HAS_NO_BYTES,
    NO_READER_FOR_KIND,
    Document,
    FileKind,
)
from src.infrastructure.excel import Template
from src.infrastructure.storage.workbooks import WorkbookStore
from src.services.extraction import DocumentRole, ReadDocument, RfqExtraction
from src.services.extraction.models import (
    DATE_AMBIGUOUS,
    DATE_IMPOSSIBLE,
    IMO_INVALID,
    READ_FAILED,
    HeaderField,
    HeaderValue,
    LineItem,
    NormalizedHeader,
    RfqHeader,
)
from src.services.workbook import (
    BRANCH_CELL,
    CELLS,
    FIRST_ITEM_ROW,
    ITEMS_TRUNCATED,
    RECEIVED_CELLS,
    REMARKS_CELL,
    SUBJECT_CELL,
    Subject,
    WorkbookBuilder,
)
from tests.workbook_builder import master

RECEIVED = datetime(2026, 9, 5, 8, 14, tzinfo=UTC)


def extraction(
    *,
    text: dict[HeaderField, str] | None = None,
    dates: dict[HeaderField, date] | None = None,
    items: list[LineItem] | None = None,
    dropped: dict[HeaderField, str] | None = None,
    checked: dict[HeaderField, str] | None = None,
    documents: list[ReadDocument] | None = None,
    unmapped_columns: list[str] | None = None,
) -> RfqExtraction:
    normalized = NormalizedHeader(
        text=dict(text or {}),
        dates=dict(dates or {}),
        dropped=dict(dropped or {}),
        checked=dict(checked or {}),
    )
    header = RfqHeader(
        fields={
            name: HeaderValue(value=value, source="email.body")
            for name, value in (text or {}).items()
        }
    )
    return RfqExtraction(
        header=header,
        normalized=normalized,
        items=list(items or []),
        documents=list(documents or []),
        unmapped_columns=list(unmapped_columns or []),
    )


def attachment(
    name: str,
    *,
    role: DocumentRole = DocumentRole.EMPTY,
    warnings: tuple[str, ...] = (),
    read_warnings: tuple[str, ...] = (),
) -> ReadDocument:
    """One attachment as stage B hands it over, without stage B running."""
    return ReadDocument(
        document=Document(
            filename=name, kind=FileKind.UNKNOWN, size_bytes=2048, warnings=list(warnings)
        ),
        role=role,
        warnings=list(read_warnings),
    )


def item(number: int, **fields) -> LineItem:
    return LineItem(sr_no=number, source="Requisition.xlsx#Sheet1", **fields)


@pytest.fixture
def builder() -> WorkbookBuilder:
    return WorkbookBuilder(Template(master()))


def worksheet(filled) -> bytes:
    with zipfile.ZipFile(io.BytesIO(filled.data)) as book:
        return book.read("xl/worksheets/sheet1.xml")


def cells(filled) -> dict[str, object]:
    book = openpyxl.load_workbook(io.BytesIO(filled.data), data_only=True)
    page = book["KASS RFQ Template"]
    book.close()
    return page


# --- the form -------------------------------------------------------------


async def test_each_header_field_lands_in_the_cell_the_template_gives_it(builder):
    filled = await builder.build(
        extraction(
            text={
                HeaderField.VESSEL_NAME: "MV ALMI GLOBE",
                HeaderField.IMO: "9232395",
                HeaderField.RFQ_REFERENCE: "78432",
                HeaderField.DELIVERY_PORT: "UAE - JEBEL ALI",
                HeaderField.CURRENCY: "USD",
                HeaderField.RFQ_TYPE: "DECK",
            },
            dates={HeaderField.ETA: date(2026, 10, 12)},
        )
    )
    page = cells(filled)

    assert page["C2"].value == "MV ALMI GLOBE"
    assert page["C3"].value == 9232395, "the desk's own forms hold the IMO as a number"
    assert page["C4"].value == "78432"
    assert page["H2"].value == "UAE - JEBEL ALI"
    assert page["H3"].value == datetime(2026, 10, 12)
    assert page["H8"].value == "USD"
    assert page["H11"].value == "DECK"


async def test_a_field_nobody_found_leaves_the_cell_empty(builder):
    """The master ships `AED` in the currency cell and a Dubai port in `H2`.
    Inheriting either of those is inventing the customer's answer for them."""
    filled = await builder.build(extraction(text={HeaderField.VESSEL_NAME: "MV ALMI GLOBE"}))
    page = cells(filled)

    assert page["H8"].value is None
    assert page["H2"].value is None


async def test_the_branch_comes_from_the_routing_and_not_from_the_email(builder):
    """`C5` decides which port list `H2` validates against, so it has to be the
    branch that is about to receive the email."""
    filled = await builder.build(extraction(), branch="UAE-DUBAI")

    assert cells(filled)["C5"].value == "UAE-DUBAI"


async def test_the_subject_says_this_is_a_request_for_quote(builder):
    filled = await builder.build(extraction())

    assert cells(filled)["C10"].value == Subject.REQUEST_FOR_QUOTE.value


async def test_the_subject_is_a_choice_rather_than_a_constant(builder):
    """The workbook offers two today and the team expects more."""
    filled = await builder.build(extraction(), subject=Subject.ORDER)

    assert cells(filled)["C10"].value == "ORDER"


async def test_the_sender_code_goes_in_when_the_lookup_found_one(builder):
    """Looked up, never read - but it is a starred cell and it does get filled."""
    filled = await builder.build(extraction(text={HeaderField.SENDER_CODE: "GM8620"}))

    assert cells(filled)["H9"].value == "GM8620"


async def test_when_the_email_arrived_is_recorded_in_both_cells(builder):
    """In the master `H10` is a formula over `H12`; in the form the desk sends
    they are two timestamps holding the same moment, so both are written."""
    filled = await builder.build(extraction(), received_at=RECEIVED)
    page = cells(filled)

    assert page["H10"].value == datetime(2026, 9, 5, 8, 14)
    assert page["H12"].value == datetime(2026, 9, 5, 8, 14)


def test_the_customer_name_cell_is_not_on_the_map():
    """`A1` is a formula in the master and, in the form the desk sends, the text
    that formula falls back to. Either way it is not ours to write."""
    written = set(CELLS.values()) | {BRANCH_CELL, SUBJECT_CELL, REMARKS_CELL, *RECEIVED_CELLS}

    assert "A1" not in written


async def test_the_customer_name_formula_survives_a_template_that_has_one(builder):
    """The form we ship has no formulas left in it, but the master does, and the
    writer must not be the reason one disappears."""
    filled = await builder.build(extraction(), received_at=RECEIVED)

    assert b'<f t="array" ref="A1">' in worksheet(filled)


# --- the items ------------------------------------------------------------


async def test_the_items_go_down_in_the_order_the_customer_wrote_them(builder):
    filled = await builder.build(
        extraction(
            items=[
                item(1, description="ROPE PP 24MM", quantity="2", uom="COIL"),
                item(2, description="SHACKLE 12MM", quantity="10", uom="PCS"),
                item(3, description="PAINT PRIMER", quantity="4", uom="CAN"),
            ]
        )
    )
    page = cells(filled)

    assert [page[f"D{FIRST_ITEM_ROW + n}"].value for n in range(3)] == [
        "ROPE PP 24MM",
        "SHACKLE 12MM",
        "PAINT PRIMER",
    ]
    assert [page[f"A{FIRST_ITEM_ROW + n}"].value for n in range(3)] == [1, 2, 3]
    assert filled.items == 3


async def test_a_count_goes_in_as_a_number_so_the_sheet_can_add_it_up(builder):
    filled = await builder.build(extraction(items=[item(1, description="ROPE", quantity="12")]))

    assert cells(filled)[f"E{FIRST_ITEM_ROW}"].value == 12


async def test_a_quantity_that_is_not_plainly_a_number_stays_as_it_was_written(builder):
    """`1,5` is one and a half in half of Europe and one thousand five hundred
    in the other half. Converting it would be guessing."""
    filled = await builder.build(extraction(items=[item(1, description="ROPE", quantity="1,5")]))

    assert cells(filled)[f"E{FIRST_ITEM_ROW}"].value == "1,5"


async def test_a_field_the_customer_left_out_leaves_its_cell_empty(builder):
    filled = await builder.build(extraction(items=[item(1, description="ROPE PP 24MM")]))
    page = cells(filled)

    assert page[f"B{FIRST_ITEM_ROW}"].value is None
    assert page[f"F{FIRST_ITEM_ROW}"].value is None


async def test_the_columns_this_task_must_not_fill_stay_empty(builder):
    """Internal item code, prices and supplier belong to a later matching step."""
    filled = await builder.build(
        extraction(items=[item(1, customer_item_code="550101", description="ROPE", quantity="2")])
    )
    page = cells(filled)

    assert page[f"C{FIRST_ITEM_ROW}"].value is None
    assert [page[f"{column}{FIRST_ITEM_ROW}"].value for column in "GHIJ"] == [None] * 4


async def test_more_items_than_the_template_holds_are_reported_not_dropped_quietly(
    builder, monkeypatch
):
    monkeypatch.setattr("src.services.workbook.LAST_ITEM_ROW", FIRST_ITEM_ROW + 2)

    filled = await builder.build(
        extraction(items=[item(n, description=f"item {n}") for n in range(1, 6)])
    )

    assert filled.items == 3
    assert ITEMS_TRUNCATED in filled.warnings


# --- what comes back ------------------------------------------------------


async def test_the_file_is_named_after_what_a_person_looks_for(builder):
    filled = await builder.build(
        extraction(
            text={
                HeaderField.VESSEL_NAME: "MV ALMI GLOBE",
                HeaderField.RFQ_REFERENCE: "78432",
            }
        )
    )

    assert filled.filename == "78432.xlsx"


async def test_a_reference_that_would_break_a_filename_is_cleaned_up(builder):
    filled = await builder.build(
        extraction(text={HeaderField.RFQ_REFERENCE: "RFQ/2026\\78432: deck"})
    )

    assert filled.filename == "RFQ 2026 78432 deck.xlsx"


async def test_an_rfq_with_no_reference_of_its_own_gets_one(builder):
    """A starred cell cannot be left blank, and the file has to be called
    something. Recognisably ours, so nobody reads it as the customer's."""
    filled = await builder.build(extraction(), received_at=RECEIVED, keep_as="3f2a9c-dead")
    page = cells(filled)

    assert page["C4"].value == "POC-20260905-3F2A9C"
    assert filled.filename == "POC-20260905-3F2A9C.xlsx"
    assert "generated" in page[REMARKS_CELL].value


async def test_a_reference_the_customer_gave_is_never_replaced(builder):
    filled = await builder.build(
        extraction(text={HeaderField.RFQ_REFERENCE: "DANT260189"}), received_at=RECEIVED
    )

    assert cells(filled)["C4"].value == "DANT260189"
    assert "generated" not in (cells(filled)[REMARKS_CELL].value or "")


async def test_a_copy_missing_a_starred_field_is_not_complete(builder):
    filled = await builder.build(extraction(items=[item(1, description="ROPE")]))

    assert not filled.is_complete
    assert "vessel_name" in filled.missing_required


async def test_the_journal_line_says_which_master_the_copy_was_cut_from(builder):
    filled = await builder.build(extraction(items=[item(1, description="ROPE")]))

    payload = filled.journal_payload()

    assert payload["items"] == 1
    assert payload["size_bytes"] == len(filled.data)
    assert len(payload["template_sha256"]) == 64


# --- keeping a copy -------------------------------------------------------


async def test_the_copy_is_kept_under_the_id_of_the_decision_that_made_it(tmp_path: Path):
    store = WorkbookStore(tmp_path, enabled=True)
    builder = WorkbookBuilder(Template(master()), store)

    filled = await builder.build(extraction(), keep_as="abc-123")

    assert filled.saved_to == tmp_path / "abc-123.xlsx"
    assert filled.saved_to.read_bytes() == filled.data


async def test_nothing_is_written_when_the_store_is_switched_off(tmp_path: Path):
    builder = WorkbookBuilder(Template(master()), WorkbookStore(tmp_path, enabled=False))

    filled = await builder.build(extraction(), keep_as="abc-123")

    assert filled.saved_to is None
    assert list(tmp_path.iterdir()) == []


async def test_an_id_that_would_escape_the_folder_cannot(tmp_path: Path):
    store = WorkbookStore(tmp_path, enabled=True)

    saved = await store.save("../../etc/passwd", b"x")

    assert saved is not None
    assert saved.parent == tmp_path


# --- what a person is told on the form ------------------------------------


async def test_the_remarks_say_a_machine_filled_this_in(builder):
    """First thing a person reads. The desk has to know not to trust it blind."""
    filled = await builder.build(extraction(items=[item(1, description="ROPE")]))

    assert "Filled automatically" in cells(filled)[REMARKS_CELL].value


async def test_the_remarks_name_every_starred_cell_left_blank(builder):
    filled = await builder.build(extraction(text={HeaderField.VESSEL_NAME: "MV ALMI GLOBE"}))
    remarks = cells(filled)[REMARKS_CELL].value

    assert "Not filled in:" in remarks
    for name in ("IMO Number", "RFQ Reference", "Delivery Port", "Currency", "RFQ Type"):
        assert name in remarks
    assert "Vessel Name" not in remarks


async def test_the_remarks_carry_the_models_account_of_what_was_not_there(builder):
    """Two kinds of reason, and they answer different questions. Code answers
    for a value we refused; the model answers for a value that is simply not in
    the email, because it is the only thing that read the email."""
    header = RfqHeader(
        not_found={HeaderField.IMO: "Not stated anywhere - the vessel is named but never numbered."}
    )
    filled = await builder.build(
        RfqExtraction(
            header=header,
            normalized=NormalizedHeader(dropped={HeaderField.ETA: DATE_AMBIGUOUS}),
            items=[item(1, description="ROPE")],
        )
    )
    remarks = cells(filled)[REMARKS_CELL].value

    assert "IMO Number: Not stated anywhere" in remarks
    assert "E.T.A: could be read two ways" in remarks


async def test_the_email_and_the_form_say_the_same_sentences(builder):
    """The desk reads one of the two. They must not differ."""
    header = RfqHeader(
        not_found={HeaderField.CURRENCY: "Neither the email nor the requisition names one."}
    )
    filled = await builder.build(
        RfqExtraction(
            header=header,
            normalized=NormalizedHeader(),
            items=[item(1, description="ROPE")],
        )
    )
    remarks = cells(filled)[REMARKS_CELL].value

    for line in filled.why_blank:
        assert line in remarks
    assert "Currency: Neither the email nor the requisition names one." in filled.why_blank


async def test_the_remarks_give_the_reason_when_there_is_one(builder):
    """A cell nobody could fill and a cell whose value we refused look the same
    on the form, and only one of them has anything to explain."""
    filled = await builder.build(extraction(dropped={HeaderField.ETA: DATE_AMBIGUOUS}))

    assert "E.T.A: could be read two ways" in cells(filled)[REMARKS_CELL].value


async def test_the_remarks_tell_a_customers_typo_from_our_own_limit(builder):
    """`31/02` reads perfectly and is not a day. The person finishing this form
    has to know whether to open the customer's file or to report a parser."""
    filled = await builder.build(extraction(dropped={HeaderField.ETA: DATE_IMPOSSIBLE}))

    assert "E.T.A: no such day in the calendar" in cells(filled)[REMARKS_CELL].value


async def test_the_remarks_say_an_imo_did_not_add_up(builder):
    """Written all the same - a number with a typo still says which vessel."""
    filled = await builder.build(
        extraction(text={HeaderField.IMO: "9425712"}, checked={HeaderField.IMO: IMO_INVALID})
    )
    page = cells(filled)

    assert page["C3"].value == 9425712
    assert "check digit does not add up" in page[REMARKS_CELL].value


async def test_the_remarks_name_a_file_nobody_could_read(builder):
    """Measured on a live run: two files arrived that were never read - one in a
    format we cannot open, one whose bytes Graph would not hand over - and the
    form said nothing at all about either. A silently dropped requisition is
    the worst thing this agent can do."""
    filled = await builder.build(
        extraction(
            items=[item(1, description="ROPE")],
            documents=[
                attachment("Specification.doc", warnings=(NO_READER_FOR_KIND,)),
                attachment("FW enquiry.eml", warnings=(ATTACHMENT_HAS_NO_BYTES,)),
            ],
        )
    )
    remarks = cells(filled)[REMARKS_CELL].value

    assert "Specification.doc: a format this agent cannot open" in remarks
    assert "FW enquiry.eml: the mail server did not hand over its contents" in remarks


async def test_the_forwarded_email_names_them_as_well(builder):
    """The desk reads the email before it opens the form."""
    filled = await builder.build(
        extraction(
            items=[item(1, description="ROPE")],
            documents=[attachment("Scan.pdf", read_warnings=(READ_FAILED,))],
        )
    )

    assert filled.unread == ["Scan.pdf: the reader could not answer about it"]


async def test_a_zip_that_gave_up_its_contents_is_not_reported_as_unread(builder):
    """A container has nothing of its own to show, and that is how it works.
    Reporting it would teach the desk to ignore this section."""
    filled = await builder.build(
        extraction(
            items=[item(1, description="ROPE")],
            documents=[
                attachment("bundle.zip"),
                attachment("bundle.zip > items.csv", role=DocumentRole.ITEM_GRID),
            ],
        )
    )

    assert filled.unread == []
    assert "bundle.zip" not in (cells(filled)[REMARKS_CELL].value or "")


async def test_a_photo_that_showed_something_is_not_reported_as_unread(builder):
    """It carried no items and was still read: its identifiers are in the
    journal and the header saw them."""
    read = attachment("nameplate.jpg", role=DocumentRole.SUPPORTING)
    filled = await builder.build(
        extraction(
            items=[item(1, description="ROPE")],
            documents=[ReadDocument(document=read.document, role=read.role, facts=["ICOM"])],
        )
    )

    assert filled.unread == []


async def test_the_remarks_name_the_columns_that_went_nowhere(builder):
    """The customer wrote "urgent" against three rows and no cell of this form
    takes it. Measured: it was in the journal and nowhere a person would see."""
    filled = await builder.build(
        extraction(
            items=[item(1, description="ROPE")],
            unmapped_columns=["F (REMARKS)"],
        )
    )
    remarks = cells(filled)[REMARKS_CELL].value

    assert "Columns of the customer's table" in remarks
    assert "F (REMARKS)" in remarks


async def test_the_remarks_say_when_no_items_could_be_read(builder):
    filled = await builder.build(extraction())

    assert "No line items could be read" in cells(filled)[REMARKS_CELL].value


async def test_the_remarks_go_under_the_form_and_not_into_a_value_cell(builder):
    """Prose in `H3` would be read as a date by whatever opens the form next."""
    filled = await builder.build(extraction(dropped={HeaderField.ETA: DATE_AMBIGUOUS}))
    page = cells(filled)

    assert page["H3"].value is None
    assert page["A14"].value is not None


# --- the unit of measure --------------------------------------------------


async def test_the_unit_goes_in_as_the_customer_wrote_it_in_capitals(builder):
    """Not mapped through the workbook's own table: it reads `BOX -> NULL` and
    `BTL -> UOM`, and the desk's finished RFQs carry `PC` beside `PCS`."""
    filled = await builder.build(
        extraction(items=[item(1, description="ROPE", uom="coil"),
                          item(2, description="PAINT", uom="Pcs")])
    )
    page = cells(filled)

    assert page[f"F{FIRST_ITEM_ROW}"].value == "COIL"
    assert page[f"F{FIRST_ITEM_ROW + 1}"].value == "PCS"


# --- what goes in as a number ---------------------------------------------


async def test_a_customer_item_code_of_digits_goes_in_as_a_number(builder):
    """Measured: `B` is a numeric column in both of the desk's finished RFQs."""
    filled = await builder.build(
        extraction(items=[item(1, customer_item_code="696737", description="U-BOLT")])
    )

    assert cells(filled)[f"B{FIRST_ITEM_ROW}"].value == 696737


async def test_a_code_padded_with_zeros_keeps_them(builder):
    """`0012345` is a code somebody's system padded, not twelve thousand and
    something - and a number cell cannot hold the difference."""
    filled = await builder.build(
        extraction(items=[item(1, customer_item_code="0012345", description="ROPE")])
    )

    assert cells(filled)[f"B{FIRST_ITEM_ROW}"].value == "0012345"


# --- the zone the timestamps are written in --------------------------------


async def test_the_timestamp_is_written_in_utc_by_default():
    """A cell carries no zone of its own, so this is a decision, not a detail."""
    from datetime import timedelta, timezone

    builder = WorkbookBuilder(Template(master()))
    singapore = datetime(2026, 9, 5, 16, 14, tzinfo=timezone(timedelta(hours=8)))

    filled = await builder.build(extraction(), received_at=singapore)

    assert cells(filled)["H12"].value == datetime(2026, 9, 5, 8, 14)


async def test_a_desk_can_have_its_own_zone():
    from zoneinfo import ZoneInfo

    builder = WorkbookBuilder(Template(master()), timezone=ZoneInfo("Asia/Singapore"))

    filled = await builder.build(extraction(), received_at=RECEIVED)

    assert cells(filled)["H12"].value == datetime(2026, 9, 5, 16, 14)


async def test_the_remarks_can_be_switched_off():
    """The desk's own finished RFQs leave the block empty, so this is theirs to
    decide - the forwarded email and the journal say it either way."""
    builder = WorkbookBuilder(Template(master()), remarks=False)

    filled = await builder.build(extraction())

    assert cells(filled)[REMARKS_CELL].value is None
